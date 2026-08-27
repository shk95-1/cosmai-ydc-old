"""산출물 자기 검사. 사람이 26장을 눈으로 훑는 대신 규칙으로 잡는다."""
import re
from pptx import Presentation

P = r'C:\Users\Admin\Downloads\[1팀_리메릭]발표자료_260827_COSMAI.pptx'
prs = Presentation(P)
BAD = ["예)", "※", "[이름]", "[담당", "[핵심 지표", "[효율", "[비용", "[시간",
       "0.0억", "00시간", "[단계", "[지표", "[문제", "[영향",
       "작성 항목", "이 영역에 배치", "예: ", "OO", "YOLO", "mAP",
       "프로젝트 발표자료", "[팀명", "[프로젝트"]
BANNED = ["주력", "양쪽 다 진다", "결함이 아니다", "잘 되는", "훌륭", "우수한"]

problems, sections, pages = [], [], []
for i, s in enumerate(prs.slides, 1):
    texts = []
    for sh in s.shapes:
        if sh.has_text_frame:
            texts.append((sh, sh.text_frame.text))
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    texts.append((sh, c.text))
    joined = "\n".join(t for _, t in texts)
    for b in BAD:
        if b in joined:
            problems.append(f"[{i}] 양식 잔재 {b!r}")
    if re.search(r"(?<!\d)00%", joined):
        problems.append(f"[{i}] 양식 잔재 '00%'")
    for b in BANNED:
        if b in joined:
            problems.append(f"[{i}] 금지 표현 {b!r}")
    # 섹션 라벨 · 페이지 번호
    for sh, t in texts:
        if sh.top is None:
            continue
        top = sh.top / 914400
        if abs(top - 0.40) < 0.06 and sh.left / 914400 < 1.0 and t.strip():
            sections.append((i, t.strip()))
        if abs(top - 7.05) < 0.06 and re.fullmatch(r"\d{1,2}", t.strip()):
            pages.append((i, int(t.strip())))
    # 슬라이드 밖으로 나간 도형
    for sh in s.shapes:
        if sh.left is None or sh.width is None:
            continue
        r = (sh.left + sh.width) / 914400
        b = (sh.top + sh.height) / 914400
        if r > 13.40 or b > 7.55 or sh.left / 914400 < -0.02:
            problems.append(f"[{i}] {sh.name} 화면 밖 (right={r:.2f} bottom={b:.2f})")

print(f"총 {len(prs.slides)}장\n")
print("=== 섹션 흐름 ===")
for i, t in sections:
    print(f"  {i:>2}  {t}")
bad_pages = [(i, p) for i, p in pages if i != p]
print(f"\n페이지 번호 슬롯 {len(pages)}개 · 불일치 {len(bad_pages)}  {bad_pages}")
print(f"\n=== 문제 {len(problems)}건 ===")
for p in problems:
    print("  " + p)
