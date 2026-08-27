"""추가 슬라이드 4장 — 데이터 인제션 2장 + LLM 극성(감성) 2장.

    python deck/build_add.py

**별도 파일로 낸다.** 시현님이 직접 고친 30장 본편을 덮어쓰지 않는다.
파워포인트에서 슬라이드를 복사해 아래 자리에 끼우면 된다.

    1·2장 → 03 솔루션 설계, 15p 「전체 파이프라인」 바로 뒤
    3·4장 → 04 구현·검증, 17p 「주요 기능 및 구현 결과」 바로 앞

출처는 강석현 tech-talk (`cosmai-tech-talk.html`) 이고, 커머스·DB·LLM 극성은
석현님 담당이라 대본에서 「제가」로 말할 수 있는 장이다.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from pptlib import (GREEN, INK, DIM, NOTE, WHITE, S1, S2, S3,
                    set_text, blank_from, renumber, move_slide,
                    hbar, legend, panel, arrow, _txt, _bar)

TPL = Path(r"C:\Users\Admin\Downloads\NVIDIA규격_PoC_프로젝트_발표자료_양식.pptx")
OUT = Path(r"C:\Users\Admin\Downloads\추가슬라이드_인제션_감성분석.pptx")

prs = Presentation(TPL)
base = len(prs.slides)


def band(sl, x, y, w, chip, chip_col, items, h=0.62, item_h=0.38, size=8.5):
    """가로 밴드 한 줄 — 왼쪽 칩 + 오른쪽 항목 박스들."""
    panel(sl, x, y, w, h, fill="FAF9FC", line="E2DEEA", radius=0.12)
    _bar(sl, x + 0.10, y + (h - 0.34) / 2, 0.86, 0.34, chip_col)
    _txt(sl, x + 0.10, y + (h - 0.34) / 2, 0.86, 0.34, chip, size=9.5,
         color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    n = len(items)
    iw = (w - 1.18) / n - 0.08
    for i, it in enumerate(items):
        ix = x + 1.08 + i * (iw + 0.08)
        panel(sl, ix, y + (h - item_h) / 2, iw, item_h, fill="FFFFFF",
              line="DCD8E4", radius=0.14)
        _txt(sl, ix + 0.06, y + (h - item_h) / 2, iw - 0.12, item_h, it,
             size=size, color=INK, align=PP_ALIGN.CENTER)
    return y + h


def kv(sl, x, y, w, rows, kw=2.35, row_h=0.42, size=9.5):
    """왼쪽 항목 / 오른쪽 설명 — 표 없이 줄무늬로."""
    for i, (k, v) in enumerate(rows):
        ry = y + i * row_h
        if i % 2 == 0:
            panel(sl, x, ry, w, row_h - 0.04, fill="FAF9FC", line="FAF9FC")
        _bar(sl, x + 0.10, ry + row_h / 2 - 0.09, 0.06, 0.16, S1)
        _txt(sl, x + 0.24, ry, kw - 0.30, row_h - 0.04, k, size=size,
             color=INK, bold=True)
        _txt(sl, x + kw, ry, w - kw - 0.14, row_h - 0.04, v, size=size, color=DIM)
    return y + len(rows) * row_h


# ══════════════════════════════════════════════════════════════
# 1  데이터 인제션 — 버리지 않고 표시한다
# ══════════════════════════════════════════════════════════════
s1 = blank_from(prs, 2, base, "03  솔루션 설계",
                "데이터 인제션 — 버리지 않고 표시한다",
                "수집·정제 단계를 관통하는 원칙은 하나입니다. 지우지 않고 플래그로 남깁니다.")

panel(s1, 0.50, 2.20, 12.33, 0.72, fill="F2EFFA", line="D9D2F0", radius=0.10)
_txt(s1, 0.72, 2.20, 11.90, 0.72,
     ["빈 텍스트 · 중복 댓글 · 광고 영상 · 부분 수집을 지우지 않고 플래그로 남깁니다.",
      "그래야 「빼면 결론이 어떻게 달라지나」를 나중에 물을 수 있습니다 — 지운 데이터로는 그 질문을 못 합니다."],
     size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE)

y = 3.10
y = band(s1, 0.50, y, 12.33, "① 수집", S1,
         ["전수 · 상한 없이", "원본 JSONL 그대로 보존", "그 옆에 정규화 CSV",
          "매니페스트에 호출 수 · 쿼터 · 한계"]) + 0.16
arrow(s1, 6.61, y - 0.14, 0.22, 0.15)
y += 0.12
y = band(s1, 0.50, y, 12.33, "② 정제", S3,
         ["HTML 엔티티", "→ NFKC 정규화", "→ 제어문자 제거", "→ 공백 축약"]) + 0.16
arrow(s1, 6.61, y - 0.14, 0.22, 0.15)
y += 0.12
y = band(s1, 0.50, y, 12.33, "③ 플래그", S2,
         ["빈 텍스트", "중복 댓글", "광고 · 협찬", "부분 수집"])

_txt(s1, 0.50, y + 0.28, 6.05, 1.00,
     ["정규화는 함수 하나만 씁니다",
      "소스마다 정규화가 다르면 소스 간 비교가 무의미해집니다. "
      "네 소스가 같은 함수를 통과하도록 강제했습니다."],
     size=10, color=INK, anchor=MSO_ANCHOR.TOP)
_txt(s1, 6.78, y + 0.28, 6.05, 1.00,
     ["상한을 없앤 이유",
      "1차 수집은 채널당 상한 10이었고, 그 값이 「매달 10건 트렌드」로 "
      "보였습니다. 지금은 43채널 13,979편을 전수로 받습니다."],
     size=10, color=INK, anchor=MSO_ANCHOR.TOP)

_txt(s1, 0.50, y + 1.16, 12.33, 0.30,
     "댓글 60,348건 · 중복 제외 60,311건 — 이 차이도 지우지 않고 플래그로 남긴 결과입니다.",
     size=9.5, color=NOTE)
print("1  데이터 인제션")


# ══════════════════════════════════════════════════════════════
# 2  공통 스키마 5계층과 운영
# ══════════════════════════════════════════════════════════════
s2 = blank_from(prs, 2, base + 1, "03  솔루션 설계",
                "공통 스키마 5계층 — 네 레포의 유일한 결합 지점",
                "레포는 사람 수만큼 나뉘어 있지만 붙는 곳은 한 군데입니다. document 테이블입니다.")

LAY = [("raw_record", "받은 그대로", ""), ("document", "정규화 본문", "261,317행"),
       ("mention", "주제·성분 언급", "105,358행"), ("entity_link", "성분 표준화", "INCI 매핑"),
       ("topic_daily_metric", "구성비 지표", "")]
lw, gap = 2.28, 0.24
for i, (name, desc, cnt) in enumerate(LAY):
    lx = 0.50 + i * (lw + gap)
    col = S1 if i in (0, 1) else (S3 if i < 4 else S2)
    panel(s2, lx, 2.28, lw, 1.24, fill="FFFFFF", line="DCD8E4", radius=0.06)
    _bar(s2, lx, 2.28, lw, 0.09, col)
    _txt(s2, lx + 0.10, 2.46, lw - 0.20, 0.34, name, size=10.5, color=INK,
         bold=True, align=PP_ALIGN.CENTER)
    _txt(s2, lx + 0.10, 2.80, lw - 0.20, 0.28, desc, size=9, color=DIM,
         align=PP_ALIGN.CENTER)
    if cnt:
        _txt(s2, lx + 0.10, 3.10, lw - 0.20, 0.30, cnt, size=10.5, color=GREEN,
             bold=True, align=PP_ALIGN.CENTER)
    if i < len(LAY) - 1:
        arrow(s2, lx + lw + 0.02, 2.83, 0.20, 0.15)

_txt(s2, 0.50, 3.66, 12.33, 0.28,
     "네 소스가 각자 다르게 수집·정제해도 document 에서 만나 같은 계약을 지킵니다. "
     "그 아래는 소스별이고, 그 위는 공통입니다.", size=10, color=INK)

panel(s2, 0.50, 4.08, 6.05, 2.42, fill="FAF9FC", line="E2DEEA", radius=0.03)
_txt(s2, 0.72, 4.18, 5.6, 0.28, "운영 — 한 소스가 죽어도 나머지가 돈다",
     size=11, color=S1, bold=True)
kv(s2, 0.62, 4.52, 5.80, [
    ("supercronic × 6", "컨테이너 = 스케줄 단위. TZ 고정"),
    ("PostgREST", "적재 후 조회는 REST 로. 별도 API 서버 없음"),
    ("정적 JS 포털", "프레임워크 없는 6화면. node --test 를 push 게이트에"),
    ("OpenSSL 3.5+", "bookworm TLS 는 oliveyoung Cloudflare 챌린지에 막힘"),
], kw=1.75, row_h=0.46, size=9)

panel(s2, 6.78, 4.08, 6.05, 2.42, fill="FAF9FC", line="E2DEEA", radius=0.03)
_txt(s2, 7.00, 4.18, 5.6, 0.28, "계약 — 기계가 검사할 수 있는 것만",
     size=11, color=S3, bold=True)
kv(s2, 6.90, 4.52, 5.80, [
    ("DDL 은 추가만", "DROP · 타입 변경 · 데이터 손실은 거부. 정규식 화이트리스트"),
    ("3롤 멱등 bootstrap", "몇 번 돌려도 같은 상태. DDL 거부 음성 테스트로 확인"),
    ("진짜 Postgres 로 테스트", "테스트당 스키마 · 일회용 컨테이너. SQLite 로 안 한다"),
    ("비밀은 트리 밖", "레포에는 이름만. 네 레포 이력에 자격증명 0건"),
], kw=1.95, row_h=0.46, size=9)

_txt(s2, 0.50, 6.62, 12.33, 0.30,
     "「프로덕션은 Postgres, 테스트는 SQLite」가 방언 버그를 내보냅니다. "
     "그래서 테스트도 진짜 Postgres 를 띄웁니다 — 컨테이너는 테스트가 끝나면 버립니다.",
     size=9.5, color=NOTE)
print("2  공통 스키마 5계층")


# ══════════════════════════════════════════════════════════════
# 3  LLM 극성 판정 — LLM 은 한 지점에만
# ══════════════════════════════════════════════════════════════
s3 = blank_from(prs, 2, base + 2, "04  구현 · 검증",
                "리뷰 감성 분석 — LLM 은 한 지점에만 둔다",
                "규칙이 기본이고 LLM 은 예외입니다. 규칙이 못 넘는 자리에서만 부릅니다.")

STAGE = [
    ("① 규칙", S1, "사전 · 패턴 기반 극성 판정",
     ["설치·의존성 없음", "결과가 결정적이라 재현됨", "비용 0"]),
    ("② 게이트", S3, "400문장 평가셋을 넘느냐",
     ["규칙이 넘으면 여기서 끝", "못 넘는 구간만 다음으로", "판정은 같은 입력·같은 출력 비교로만"]),
    ("③ LLM", S2, "claude-sonnet-5 · structured output",
     ["Batches API 로 묶어 호출", "$10 하드스톱", "호출 전 예산 예약 후 실측으로 덮어씀"]),
]
cw = 3.94
for i, (name, col, sub, bullets) in enumerate(STAGE):
    cx = 0.50 + i * (cw + 0.28)
    panel(s3, cx, 2.28, cw, 1.82, fill="FFFFFF", line="DCD8E4", radius=0.04)
    _bar(s3, cx, 2.28, cw, 0.09, col)
    _txt(s3, cx + 0.22, 2.48, cw - 0.44, 0.32, name, size=13, color=INK, bold=True)
    _txt(s3, cx + 0.22, 2.82, cw - 0.44, 0.30, sub, size=9.5, color=col, bold=True)
    _txt(s3, cx + 0.22, 3.18, cw - 0.44, 1.24,
         ["· " + b for b in bullets], size=9.5, color=DIM, anchor=MSO_ANCHOR.TOP)
    if i < 2:
        arrow(s3, cx + cw + 0.03, 3.32, 0.22, 0.16)

panel(s3, 0.50, 4.32, 12.33, 1.62, fill="F2EFFA", line="D9D2F0", radius=0.04)
_txt(s3, 0.72, 4.42, 11.90, 0.30,
     "왜 이렇게 나눴나 — 세 가지 이유입니다", size=11, color=INK, bold=True)
kv(s3, 0.62, 4.76, 12.10, [
    ("재현성", "LLM 출력은 같은 입력에도 흔들립니다. 규칙이 낼 수 있는 구간을 LLM 에 맡기면 재현 가능한 부분을 잃습니다"),
    ("비용 상한", "$10 하드스톱을 호출 전에 예약합니다. 예산을 다 쓰면 판정을 멈추고, 어디까지 했는지 남깁니다"),
    ("교체 가능성", "같은 프롬프트를 ollama 로컬 모델로도 왕복시켜 비교합니다. 한 벤더에 묶이지 않기 위한 것입니다"),
], kw=1.55, row_h=0.40, size=9.5)

_txt(s3, 0.50, 6.10, 12.33, 0.30,
     "네 레포 전체에서 LLM 이 있는 곳은 두 군데뿐입니다 — 여기 리뷰 극성 판정과, 질의 GUI 의 도구 선택입니다.",
     size=9.5, color=NOTE)
print("3  리뷰 감성 분석")


# ══════════════════════════════════════════════════════════════
# 4  극성 교체를 어떻게 판정하나
# ══════════════════════════════════════════════════════════════
s4 = blank_from(prs, 2, base + 3, "04  구현 · 검증",
                "규칙에서 LLM 으로 바꿀지, 무엇을 보고 판정했나",
                "「LLM 이 더 잘한다」는 느낌으로 바꾸지 않았습니다. 같은 입력에 같은 출력을 비교했습니다.")

panel(s4, 0.50, 2.28, 6.05, 2.26, fill="FAF9FC", line="E2DEEA", radius=0.03)
_txt(s4, 0.72, 2.38, 5.6, 0.28, "교체 판정 절차", size=11, color=S1, bold=True)
_txt(s4, 0.72, 2.70, 5.6, 1.70,
     ["① 400문장 평가셋을 고정합니다. 문장을 바꾸면 비교가 깨집니다",
      "② 규칙으로 판정해 출력을 바이트 단위로 저장합니다 (골든)",
      "③ LLM 으로 같은 400문장을 판정합니다",
      "④ 두 출력을 대조합니다 — 사람이 눈으로 보지 않습니다",
      "⑤ LLM 이 이긴 구간만 LLM 으로 넘깁니다"],
     size=9.5, color=INK, anchor=MSO_ANCHOR.TOP)

panel(s4, 6.78, 2.28, 6.05, 2.26, fill="FAF9FC", line="E2DEEA", radius=0.03)
_txt(s4, 7.00, 2.38, 5.6, 0.28, "LLM 에 넘기지 않은 것", size=11, color=S2, bold=True)
kv(s4, 6.90, 2.70, 5.80, [
    ("숫자", "「숫자를 직접 만들지 마라」가 시스템 프롬프트 첫 줄입니다"),
    ("SQL 생성", "생성된 SQL 은 검증할 수 없습니다. 고정 쿼리 도구 8개로 대체"),
    ("집계", "숫자는 Postgres 가 만들고 LLM 은 어느 도구를 부를지만 정합니다"),
    ("쓰기", "되돌릴 수 없는 작업은 LLM 경로에 두지 않았습니다"),
], kw=1.35, row_h=0.42, size=9)

panel(s4, 0.50, 4.70, 12.33, 1.66, fill="FFFFFF", line="DCD8E4", radius=0.04)
_txt(s4, 0.72, 4.80, 11.90, 0.30,
     "감성 분석 결과를 쓸 때의 한계 — 카드에 함께 싣습니다",
     size=11, color=INK, bold=True)
kv(s4, 0.62, 5.14, 12.10, [
    ("평가셋 크기", "400문장입니다. 이 숫자로 정확도를 주장하지 않고 「규칙보다 나은 구간이 있다」까지만 말합니다"),
    ("리뷰 편향", "커머스 리뷰는 구매자만 씁니다. 안 산 사람의 극성은 이 데이터에 없습니다"),
    ("제품 바스켓", "수집 창이 바뀌면 제품 구성이 바뀝니다(48 → 25제품, 교집합 14). 순위만 쓰고 수준은 안 씁니다"),
], kw=1.75, row_h=0.42, size=9.5)

_txt(s4, 0.50, 6.50, 12.33, 0.30,
     "극성은 판정에 직접 쓰지 않고 근거로만 씁니다 — 카드의 주장은 구성비에서 나오고, 극성은 그 옆에 붙습니다.",
     size=9.5, color=NOTE)
print("4  교체 판정")


# 양식의 원본 19장을 지우고 추가 4장만 남긴다
for i in range(base - 1, -1, -1):
    sldIds = prs.slides._sldIdLst
    ids = list(sldIds)
    rid = ids[i].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    prs.part.drop_rel(rid)
    sldIds.remove(ids[i])

renumber(prs)
prs.save(OUT)
print(f"\n{len(prs.slides)}장 → {OUT}")
